from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_payload(paragraph):
    return {
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
        "runs": [
            {
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "size_pt": run.font.size.pt if run.font.size else None,
                "font": run.font.name,
            }
            for run in paragraph.runs
        ],
    }


def inspect(path: Path):
    doc = Document(path)
    result = {
        "path": str(path),
        "paragraphs": [paragraph_payload(p) for p in doc.paragraphs],
        "tables": [],
        "sections": [],
        "headers": [],
        "footers": [],
    }
    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                cells.append({
                    "row": ri,
                    "col": ci,
                    "text": "\n".join(p.text for p in cell.paragraphs),
                    "paragraphs": [paragraph_payload(p) for p in cell.paragraphs],
                })
            rows.append(cells)
        result["tables"].append({"index": ti, "rows": rows})
    for si, section in enumerate(doc.sections):
        result["sections"].append({
            "index": si,
            "width_in": section.page_width.inches,
            "height_in": section.page_height.inches,
            "top_in": section.top_margin.inches,
            "bottom_in": section.bottom_margin.inches,
            "left_in": section.left_margin.inches,
            "right_in": section.right_margin.inches,
            "header_in": section.header_distance.inches,
            "footer_in": section.footer_distance.inches,
        })
        result["headers"].append([paragraph_payload(p) for p in section.header.paragraphs])
        result["footers"].append([paragraph_payload(p) for p in section.footer.paragraphs])
    return result


if __name__ == "__main__":
    for item in sys.argv[1:]:
        path = Path(item)
        out = path.with_suffix(".structure.json")
        out.write_text(json.dumps(inspect(path), ensure_ascii=False, indent=2), encoding="utf-8")
        print(out)
